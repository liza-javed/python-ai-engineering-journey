"""Heuristic resume parser.

Extracts contact details, sections, experience, education, certifications,
projects, and skills from resume text. It works on clean, text-based PDF and
DOCX files. Scanned (image-only) PDFs are not supported and raise ValueError.
"""
from __future__ import annotations

import io
import logging
import re
from functools import lru_cache
from typing import BinaryIO

import pdfplumber
import spacy
from docx import Document

from app.config import settings
from app.schemas import (
    Certification,
    ContactInfo,
    Education,
    Experience,
    Project,
    ResumeParsed,
)
from app.services.skill_vocabulary import extract_skills

logger = logging.getLogger(__name__)

SECTION_HEADERS: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "professional summary", "profile", "objective", "about me"),
    "experience": ("experience", "work experience", "professional experience", "employment", "work history"),
    "education": ("education", "academic background", "academics"),
    "skills": ("skills", "technical skills", "core skills", "competencies"),
    "certifications": ("certifications", "certificates", "licenses", "licenses & certifications"),
    "projects": ("projects", "personal projects", "key projects", "portfolio"),
}
_HEADER_LOOKUP = {alias: section for section, aliases in SECTION_HEADERS.items() for alias in aliases}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w-]+", re.IGNORECASE)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w-]+", re.IGNORECASE)
URL_RE = re.compile(r"(?:https?://|www\.)[^\s,;)]+|github\.com/[^\s,;)]+", re.IGNORECASE)
DATE_RANGE_RE = re.compile(
    r"((?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2}|Present|Current|Now)",
    re.IGNORECASE,
)
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
DEGREE_RE = re.compile(
    r"\b(Bachelor|Master|Ph\.?\s?D|Doctorate|Associate|MBA|B\.?Sc|M\.?Sc|B\.?Eng|M\.?Eng|B\.?Tech|M\.?Tech|Diploma)\b[^\n]*",
    re.IGNORECASE,
)
INSTITUTION_RE = re.compile(
    r"((?:[A-Z][\w&.'-]*\s+){0,4}(?:University|College|Institute|School|Academy)(?:\s+of\s+[A-Z][\w ]+)?)"
)
BULLET_PREFIXES = ("-", "•", "*", "▪", "◦", "·")
ACTION_VERBS = {
    "developed", "built", "designed", "implemented", "led", "managed", "created", "improved",
    "optimized", "reduced", "increased", "delivered", "architected", "deployed", "automated",
    "mentored", "engineered", "established", "migrated", "resolved",
}


@lru_cache(maxsize=1)
def _load_spacy():
    try:
        return spacy.load(settings.spacy_model)
    except OSError as exc:
        raise RuntimeError(
            f"spaCy model '{settings.spacy_model}' is not installed. "
            f"Run: python -m spacy download {settings.spacy_model}"
        ) from exc


class ResumeParser:
    """Parses resume files or raw resume text into a ResumeParsed object."""

    def __init__(self) -> None:
        self._nlp = _load_spacy()

    # ------------------------------------------------------------------ text
    @staticmethod
    def extract_text(stream: BinaryIO, extension: str) -> str:
        """Extract plain text from a PDF or DOCX stream."""
        if extension == "pdf":
            with pdfplumber.open(stream) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif extension == "docx":
            document = Document(stream)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            text = "\n".join(parts)
        else:
            raise ValueError(f"Unsupported file type: {extension}")

        text = text.strip()
        if len(text.split()) < 20:
            raise ValueError(
                "Could not read enough text from this file. Scanned or image-only PDFs are not supported."
            )
        return text

    def parse_file(self, content: bytes, extension: str) -> tuple[ResumeParsed, str]:
        """Parse uploaded bytes. Returns the parsed resume and the raw text."""
        text = self.extract_text(io.BytesIO(content), extension)
        return self.parse_text(text), text

    # --------------------------------------------------------------- parsing
    def parse_text(self, text: str) -> ResumeParsed:
        lines = [line.strip() for line in text.splitlines()]
        sections = self._split_sections(lines)
        return ResumeParsed(
            contact_info=self._extract_contact(lines, text),
            summary=" ".join(sections.get("summary", [])) or None,
            experience=self._parse_experience(sections.get("experience", [])),
            education=self._parse_education(sections.get("education", [])),
            skills=extract_skills(text),
            certifications=self._parse_certifications(sections.get("certifications", [])),
            projects=self._parse_projects(sections.get("projects", [])),
        )

    @staticmethod
    def _split_sections(lines: list[str]) -> dict[str, list[str]]:
        """Group lines under recognised section headers. A header is a short line matching a known title."""
        sections: dict[str, list[str]] = {}
        current: str | None = None
        for line in lines:
            if not line:
                continue
            key = line.lower().strip(" :-–—")
            if len(line) <= 40 and key in _HEADER_LOOKUP:
                current = _HEADER_LOOKUP[key]
                sections.setdefault(current, [])
                continue
            if current is not None:
                sections[current].append(line)
        return sections

    def _extract_contact(self, lines: list[str], text: str) -> ContactInfo:
        top = "\n".join(lines[:15])
        email = EMAIL_RE.search(text)
        phone = PHONE_RE.search(top)
        linkedin = LINKEDIN_RE.search(text)
        github = GITHUB_RE.search(text)
        return ContactInfo(
            name=self._extract_name(lines[:5]),
            email=email.group(0) if email else None,
            phone=phone.group(0).strip() if phone else None,
            linkedin=_with_scheme(linkedin.group(0)) if linkedin else None,
            github=_with_scheme(github.group(0)) if github else None,
        )

    def _extract_name(self, top_lines: list[str]) -> str | None:
        """Use the first plausible name-like line, falling back to spaCy PERSON entities."""
        for line in top_lines:
            words = line.split()
            if (
                2 <= len(words) <= 4
                and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words)
                and not any(w.lower().strip(".:") in _HEADER_LOOKUP for w in words)
            ):
                return line.title() if line.isupper() else line
        doc = self._nlp(" ".join(top_lines))
        for ent in doc.ents:
            if ent.label_ == "PERSON":
                return ent.text
        return None

    @staticmethod
    def _split_entries(lines: list[str]) -> list[tuple[str, list[str]]]:
        """Group lines into (heading, detail_lines) blocks. Bullets are details; other lines start a new block."""
        entries: list[tuple[str, list[str]]] = []
        for line in lines:
            if line.startswith(BULLET_PREFIXES):
                if entries:
                    entries[-1][1].append(line.lstrip("".join(BULLET_PREFIXES)).strip())
            else:
                entries.append((line, []))
        return entries

    def _parse_experience(self, lines: list[str]) -> list[Experience]:
        """Each role starts with a line containing a date range, optionally preceded by the title line."""
        results: list[Experience] = []
        prev_line = ""
        current: dict | None = None
        for line in lines:
            if line.startswith(BULLET_PREFIXES):
                if current is not None:
                    current["bullets"].append(line.lstrip("".join(BULLET_PREFIXES)).strip())
                continue
            match = DATE_RANGE_RE.search(line)
            if match:
                if current is not None:
                    results.append(self._experience_from(current))
                heading = line[: match.start()].strip(" |–—-,") or prev_line
                current = {"heading": heading, "duration": match.group(0), "bullets": []}
            prev_line = line
        if current is not None:
            results.append(self._experience_from(current))
        return results

    @staticmethod
    def _experience_from(block: dict) -> Experience:
        parts = [p.strip() for p in re.split(r"\s[|–—]\s|\s-\s|,\s", block["heading"]) if p.strip()]
        position = parts[0] if parts else ""
        company = parts[1] if len(parts) > 1 else ""
        description = "\n".join(block["bullets"]) or None
        return Experience(
            company=company,
            position=position,
            duration=block["duration"],
            description=description,
            skills=extract_skills(description or ""),
        )

    @staticmethod
    def _parse_education(lines: list[str]) -> list[Education]:
        results: list[Education] = []
        for index, line in enumerate(lines):
            degree_match = DEGREE_RE.search(line)
            if not degree_match:
                continue
            degree_line = degree_match.group(0).strip()
            field_match = re.search(r"\bin\s+([A-Za-z&/ ]+?)(?=\s*[|,(]|$)", degree_line, re.IGNORECASE)
            following = " ".join(lines[index + 1 : index + 3])
            institution = INSTITUTION_RE.search(following)
            year = YEAR_RE.search(f"{line} {following}")
            results.append(
                Education(
                    institution=institution.group(1).strip() if institution else "",
                    degree=degree_line.split("|")[0].strip(),
                    field=field_match.group(1).strip() if field_match else "",
                    graduation_year=year.group(0) if year else None,
                )
            )
        return results

    @staticmethod
    def _parse_certifications(lines: list[str]) -> list[Certification]:
        results: list[Certification] = []
        for line in lines:
            text = line.lstrip("".join(BULLET_PREFIXES)).strip()
            if len(text) < 4:
                continue
            parts = [p.strip() for p in text.split("|") if p.strip()]
            year = YEAR_RE.search(text)
            issuer = next((p for p in parts[1:] if not YEAR_RE.fullmatch(p)), "")
            results.append(
                Certification(
                    name=parts[0],
                    issuer=issuer,
                    year=year.group(0) if year else None,
                )
            )
        return results

    @staticmethod
    def _parse_projects(lines: list[str]) -> list[Project]:
        results: list[Project] = []
        for heading, details in ResumeParser._split_entries(lines):
            name = re.split(r"\s[-–—|]\s|:", heading, maxsplit=1)[0].strip()
            inline_desc = re.split(r"\s[-–—|]\s|:", heading, maxsplit=1)[1].strip() if re.search(r"\s[-–—|]\s|:", heading) else ""
            description = "\n".join(part for part in [inline_desc, *details] if part) or None
            url = URL_RE.search(f"{heading} {' '.join(details)}")
            if not name:
                continue
            results.append(
                Project(
                    name=name,
                    description=description,
                    technologies=extract_skills(f"{heading} {' '.join(details)}"),
                    link=_with_scheme(url.group(0).rstrip(".")) if url else None,
                )
            )
        return results

    @staticmethod
    def count_action_verb_bullets(experience: list[Experience]) -> tuple[int, int]:
        """Return (bullets starting with an action verb, total bullets). Used by the ATS scorer."""
        total = 0
        strong = 0
        for exp in experience:
            for bullet in (exp.description or "").split("\n"):
                words = bullet.strip().split()
                if not words:
                    continue
                total += 1
                if words[0].lower().strip(".,:;") in ACTION_VERBS:
                    strong += 1
        return strong, total


def _with_scheme(url: str) -> str:
    return url if url.lower().startswith("http") else f"https://{url}"
